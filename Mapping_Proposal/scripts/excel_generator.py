"""Export into the provided Nouryon Excel layout, keeping existing formatting.

This uses openpyxl in the user's Python environment. The layout is read from
row 2 of 'Mapping Proposal', and row 1 of the other worksheets.
"""
from copy import copy
import logging
from pathlib import Path
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# Excel allows 32,767 characters per cell; keep a safety margin.
MAX_CELL_CHARS = 32000


def _text(value):
    """Convert to a string that fits in one Excel cell."""
    if value is None:
        return None
    return str(value)[:MAX_CELL_CHARS]


# Reads the column names from the template itself, so adding or removing a
# column in mapping_template.xlsx no longer requires changing this code.
def _headers(sheet, row):
    """Header names of one row, up to the last non-empty header cell."""
    values = [cell.value for cell in sheet[row]]
    while values and values[-1] in (None, ''):
        values.pop()
    return values


def _clear_old_rows(sheet, start_row):
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.value = None


def _write_records(sheet, headers, records, start_row):
    template_row = start_row
    for index, record in enumerate(records):
        row_number = start_row + index
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row_number, col_index)
            if row_number > template_row:
                model_cell = sheet.cell(template_row, col_index)
                if model_cell.has_style:
                    cell._style = copy(model_cell._style)
                cell.alignment = copy(model_cell.alignment)
            value = record.get(header)
            # Numbers stay numbers; everything else is written as plain text.
            if isinstance(value, (int, float, bool)):
                cell.value = value
            else:
                cell.value = _text(value)
                # openpyxl turns text starting with '=' into a formula. Generated text
                # is never a formula, so always store it as plain text.
                if cell.value is not None:
                    cell.data_type = 's'


def export_excel(mappings, source_inventory, source_files, target_files,
                 output_path: str | Path, template_path: str | Path,
                 historical_pairs=None, selected_historical_examples=None,
                 verified_historical_evidence=None):
    historical_pairs = historical_pairs or []
    selected_historical_examples = selected_historical_examples or []
    verified_historical_evidence = verified_historical_evidence or []
    template_path, output_path = Path(template_path), Path(output_path)
    if not template_path.exists():
        raise FileNotFoundError(f'Excel template not found: {template_path}')
    wb = load_workbook(template_path)

    # Erase the previous customer's content without discarding the workbook design.
    for sheet in wb:
        first_data_row = 3 if sheet.title == 'Mapping Proposal' else 2
        _clear_old_rows(sheet, first_data_row)

    main = wb['Mapping Proposal']
    headers = _headers(main, 2)
    records = []
    questions = []
    for i, m in enumerate(mappings, 1):
        target = m['target_path']
        source = m['source_path']
        source_detail = next((f for f in source_inventory if f['path'] == source), {})
        problems = m.get('validation_issues', [])
        summary = str(m.get('reason', ''))
        if problems:
            summary += ' | VALIDATION: ' + '; '.join(problems)
        mapping_id = f'MP-{i:03d}'
        records.append({
            'Mapping ID': mapping_id,
            'Runtime source group': 'New customer original XML',
            'Runtime source file example': (source_detail.get('files') or [''])[0],
            'Confidence %': m['confidence'],
            'Requires review': 'Yes' if m['requires_review'] else 'No',
            'Review priority': m.get('review_priority', ''),
            'Confidence reason': summary,
            'Evidence summary': m.get('evidence', ''),
            'Target segment path': target.rsplit('/', 1)[0],
            'Target segment name': target.rsplit('/', 2)[-2] if '/' in target else '',
            'Target field name': target.rsplit('/', 1)[-1],
            'Target field description': '',
            'Source segment path': source.rsplit('/', 1)[0] if source else '',
            'Source field name': source.rsplit('/', 1)[-1] if source else '',
            'Source full path': source,
            'Example source value': ' | '.join(source_detail.get('examples', [])[:3]),
            'Proposed rule type': m['rule_type'],
            'Fixed value': m.get('fixed_value', ''),
            'Transformation logic': m.get('transformation', ''),
            'Condition logic': m.get('condition', ''),
            'Lookup logic': m.get('lookup', ''),
            'Loop mapping logic': m.get('loop', ''),
            'Risk category': m.get('risk_category', ''),
        })
        if m['requires_review']:
            questions.append({
                'Question ID': f'RQ-{len(questions)+1:03d}',
                'Question priority': m.get('review_priority') or 'Medium',
                'Question type': 'Mapping review',
                'Related Mapping ID(s)': mapping_id,
                'Runtime source file(s)': ', '.join(source_detail.get('files', [])),
                'Target field(s)': target,
                'Source candidate(s)': source,
                'Question': m.get('review_question') or 'Confirm the proposed mapping.',
                'Why this question matters': summary,
                'Impacted mappings count': 1,
                'Suggested answer': '',
                'Impact if unanswered': m.get('risk_category', ''),
                'Status': 'Open',
            })
    _write_records(main, headers, records, 3)

    qsheet = wb['Review Questions']
    qheaders = _headers(qsheet, 1)
    _write_records(qsheet, qheaders, questions, 2)

    invsheet = wb['Runtime Field Inventory']
    iheaders = _headers(invsheet, 1)
    inventory_records = [{
        'Source full path': f['path'],
        'Source field name': f['field'],
        'Runtime files count': len(f['files']),
        'Occurrences': f['count'],
        'Runtime files': ', '.join(f['files']),
        'Example values': ' | '.join(f['examples']),
    } for f in source_inventory]
    _write_records(invsheet, iheaders, inventory_records, 2)

    csheet = wb['File Classification']
    cheaders = _headers(csheet, 1)
    file_records = [
        {
            'File': p.name, 'Classification': 'Runtime Original',
            'Role in execution': 'New customer source inventory',
            'Pair/group': 'Current customer', 'Comment': 'Unconverted customer input',
        }
        for p in source_files
    ]
    for pair in historical_pairs:
        group = f"{pair['company']} / {pair['pair_id']}"
        file_records.extend([
            {
                'File': pair['original_file'].name,
                'Classification': 'Historical Original',
                'Role in execution': 'Historical conversion evidence',
                'Pair/group': group,
                'Comment': 'Matched to converted file; may be sampled for AI context',
            },
            {
                'File': pair['converted_file'].name,
                'Classification': 'Historical Converted',
                'Role in execution': 'Common target inventory and historical evidence',
                'Pair/group': group,
                'Comment': 'Historical target structure; paired with original',
            },
        ])
    if not historical_pairs:
        file_records.extend({
            'File': p.name, 'Classification': 'Target Example',
            'Role in execution': 'Target inventory', 'Pair/group': 'Current',
            'Comment': 'Target example',
        } for p in target_files)
    _write_records(csheet, cheaders, file_records, 2)

    summary = wb['Summary']
    summary_items = [
        ('Aantal aangeleverde bestanden', len(source_files) + (2 * len(historical_pairs) if historical_pairs else len(target_files))),
        ('Aantal Runtime Original-bestanden', len(source_files)),
        ('Gekozen werkingsmodus', 'AI Mapping Prototype (human review)'),
        ('Aantal voorgestelde mappings', len(mappings)),
        ('Aantal high confidence mappings', sum(m['confidence'] >= 90 for m in mappings)),
        ('Aantal mappings requiring review', len(questions)),
        ('Aantal source paths discovered', len(source_inventory)),
        ('Historical pairs discovered', len(historical_pairs)),
        ('Historical pairs sent to Gemini per batch', len(selected_historical_examples)),
    ]
    for index, (label, value) in enumerate(summary_items, 2):
        summary.cell(index, 1).value = label
        summary.cell(index, 2).value = value

    # Field-level candidate matches are computed independently from the historical
    # XML files by scripts/historical_evidence.py. Equal values alone do not
    # prove that two fields should be mapped.
    reference_sheet = wb['Reference Evidence']
    reference_headers = _headers(reference_sheet, 1)
    reference_records = list(verified_historical_evidence)
    # Still show which pairs Gemini received, even when no meaningful value matched.
    with_matches = {row['Reference pair'] for row in reference_records}
    for pair in selected_historical_examples:
        pair_label = f"{pair['company']} / {pair['pair_id']}"
        if pair_label in with_matches:
            continue
        reference_records.append({
            'Reference pair': pair_label,
            'Historical Original': pair['original_filename'],
            'Historical Converted': pair['converted_filename'],
            'Evidence type': 'Context only (no informative exact match)',
            'Historical source path': '',
            'Historical target path': '',
            'Example value': '',
            'Evidence note': 'Pair was supplied to Gemini; no informative field-level exact match was found.',
            'Usefulness': 'Context only',
        })
    _write_records(reference_sheet, reference_headers, reference_records, 2)

    # Rule Types is restored from the template; assumptions and risks are not fabricated.
    rules = wb['Rule Types']
    # Rule Types contains definitions in the original workbook, restore them separately
    # from reference template at the end to prevent erasing the codebook.
    original = load_workbook(template_path, read_only=True)
    original_rules = original['Rule Types']
    for row in original_rules.iter_rows(min_row=2):
        for cell in row:
            rules.cell(cell.row, cell.column).value = cell.value
    original.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(output_path)
    except PermissionError:
        # Avoid losing the finished package if the original is open in Excel.
        backup_path = output_path.with_name(output_path.stem + '_new.xlsx')
        wb.save(backup_path)
        logger.warning('Original Excel file is locked. Saved to: %s', backup_path)
        output_path = backup_path
    return output_path
