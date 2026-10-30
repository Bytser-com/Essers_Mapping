from services import sftp_storage 
from openpyxl import load_workbook
import re

VERSION_PATTERN = r"^MProposal-.+_V(\d+)(_edited)?\.xlsx$"
SHEET_NAME = "Mapping Proposal"

def get_proposal_file(company):
    best_key = (0, False)
    best_file = None

    try:
        file_names = sftp_storage.list_dir(company)
    except FileNotFoundError:
        return None, None, None, None
    for file_name in file_names:
        match = re.search(VERSION_PATTERN, file_name)
        if match is None:
            continue
        version = int(match.group(1))
        is_edited = match.group(2) is not None
        key = (version, is_edited)
        if key > best_key:
            best_key = key
            best_file = file_name
    
    best_version, best_edited = best_key

    if best_file is None:
        return None, None, None, None

    path = f"{company}/{best_file}"
    uploaded = sftp_storage.get_modified_time(path)

    return sftp_storage.load(path), best_version, uploaded, best_edited

def get_list_from_excel(company):
    excel, version, uploaded, edited = get_proposal_file(company)
    if excel is None:
        return None, None, None, None
    wb = load_workbook(excel, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        return None, None, None, None
    ws = wb[SHEET_NAME]
    titles = []
    for cell in ws[2]:
        titles.append(cell.value)
    mappings = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        record = dict(zip(titles, row))
        if record.get("Mapping ID") is None:
            continue
        mapping = {                   
        "id": record.get("Mapping ID"),
        "bronveld": record.get("Source full path"),
        "doelveld": record.get("Target field name"),
        "zekerheid": record.get("Confidence %"),
        "opmerking": record.get("Confidence reason")
    }
        mappings.append(mapping) 
    wb.close()
    return mappings, version, uploaded, edited