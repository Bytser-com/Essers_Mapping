"""Export into the provided Nouryon Excel layout, keeping existing formatting.

This uses openpyxl in the user's Python environment. The layout is read from
row 2 of 'Mapping Proposal', and row 1 of the other worksheets.
"""
from copy import copy
from pathlib import Path
from openpyxl import load_workbook


def _text(value):
    """Prevent generated text from being interpreted as an Excel formula."""
    if value is None:
        return None
    text = str(value)
    if text.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + text
    return text[:32000]


def _clear_old_rows(sheet, start_row):
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.value = None


def _write_records(sheet, headers, records, start_row):
    template_row = start_row
    for index, record in enumerate(records):
        row_number = start_row + index
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row_number, col_index)
            if row_number > template_row:
                model_cell = sheet.cell(template_row, col_index)
                if model_cell.has_style:
                    cell._style = copy(model_cell._style)
                cell.alignment = copy(model_cell.alignment)
            value = record.get(header)
            cell.value = value if isinstance(value, (int, float, bool)) else _text(value)


def export_excel(mappings, source_inventory, source_files, target_files,
                 output_path: str | Path, template_path: str | Path):
    template_path, output_path = Path(template_path), Path(output_path)
    if not template_path.exists():
        raise FileNotFoundError(f'Excel template not found: {template_path}')
    wb = load_workbook(template_path)

    # Erase the previous customer's content without discarding the workbook design.
    for sheet in wb:
        first_data_row = 3 if sheet.title == 'Mapping Proposal' else 2
        _clear_old_rows(sheet, first_data_row)

    main = wb['Mapping Proposal']
    headers = [main.cell(2, col).value for col in range(1, 31)]
    records = []
    questions = []
    for i, m in enumerate(mappings, 1):
        target = m['target_path']
        source = m['source_path']
        source_detail = next((f for f in source_inventory if f['path'] == source), {})
        problems = m.get('validation_issues', [])
        summary = str(m.get('reason', ''))
        if problems:
            summary += ' | VALIDATION: ' + '; '.join(problems)
        mapping_id = f'MP-{i:03d}'
        records.append({
            'Mapping ID': mapping_id,
            'Runtime source group': 'Prototype XML input',
            'Runtime source file example': (source_detail.get('files') or [''])[0],
            'Confidence %': m['confidence'],
            'Requires review': 'Yes' if m['requires_review'] else 'No',
            'Review priority': m.get('review_priority', ''),
            'Confidence reason': summary,
            'Evidence summary': m.get('evidence', ''),
            'Target segment path': target.rsplit('/', 1)[0],
            'Target segment name': target.rsplit('/', 2)[-2] if '/' in target else '',
            'Target field name': target.rsplit('/', 1)[-1],
            'Target field description': '',
            'Source segment path': source.rsplit('/', 1)[0] if source else '',
            'Source field name': source.rsplit('/', 1)[-1] if source else '',
            'Source full path': source,
            'Example source value': ' | '.join(source_detail.get('examples', [])[:3]),
            'Proposed rule type': m['rule_type'],
            'Fixed value': m.get('fixed_value', ''),
            'Transformation logic': m.get('transformation', ''),
            'Condition logic': m.get('condition', ''),
            'Lookup logic': m.get('lookup', ''),
            'Loop mapping logic': m.get('loop', ''),
            'Risk category': m.get('risk_category', ''),
        })
        if m['requires_review']:
            questions.append({
                'Question ID': f'RQ-{len(questions)+1:03d}',
                'Question priority': m.get('review_priority') or 'Medium',
                'Question type': 'Mapping review',
                'Related Mapping ID(s)': mapping_id,
                'Runtime source file(s)': ', '.join(source_detail.get('files', [])),
                'Target field(s)': target,
                'Source candidate(s)': source,
                'Question': m.get('review_question') or 'Confirm the proposed mapping.',
                'Why this question matters': summary,
                'Impacted mappings count': 1,
                'Suggested answer': '',
                'Impact if unanswered': m.get('risk_category', ''),
                'Status': 'Open',
            })
    _write_records(main, headers, records, 3)

    qsheet = wb['Review Questions']
    qheaders = [qsheet.cell(1, c).value for c in range(1, 15)]
    _write_records(qsheet, qheaders, questions, 2)

    invsheet = wb['Runtime Field Inventory']
    iheaders = [invsheet.cell(1, c).value for c in range(1, 7)]
    inventory_records = [{
        'Source full path': f['path'],
        'Source field name': f['field'],
        'Runtime files count': len(f['files']),
        'Occurrences': f['count'],
        'Runtime files': ', '.join(f['files']),
        'Example values': ' | '.join(f['examples']),
    } for f in source_inventory]
    _write_records(invsheet, iheaders, inventory_records, 2)

    csheet = wb['File Classification']
    cheaders = [csheet.cell(1, c).value for c in range(1, 6)]
    file_records = [
        {'File': p.name, 'Classification': 'Runtime Original', 'Role in execution': 'Source inventory', 'Pair/group': 'Current', 'Comment': 'Current XML input'}
        for p in source_files
    ] + [
        {'File': p.name, 'Classification': 'Target Example', 'Role in execution': 'Target inventory', 'Pair/group': 'Current', 'Comment': 'Target XML example or specification'}
        for p in target_files
    ]
    _write_records(csheet, cheaders, file_records, 2)

    summary = wb['Summary']
    summary_items = [
        ('Aantal aangeleverde bestanden', len(source_files) + len(target_files)),
        ('Aantal Runtime Original-bestanden', len(source_files)),
        ('Gekozen werkingsmodus', 'AI Mapping Prototype (human review)'),
        ('Aantal voorgestelde mappings', len(mappings)),
        ('Aantal high confidence mappings', sum(m['confidence'] >= 90 for m in mappings)),
        ('Aantal mappings requiring review', len(questions)),
        ('Aantal source paths discovered', len(source_inventory)),
    ]
    for index, (label, value) in enumerate(summary_items, 2):
        summary.cell(index, 1).value = label
        summary.cell(index, 2).value = value

    # Rule Types remains intact as the codebook; it is not customer-specific data.
    # Reference Evidence and Assumptions & Risks are left blank rather than fabricated.
    rules = wb['Rule Types']
    # Rule Types contains definitions in the original workbook, restore them separately
    # from reference template at the end to prevent erasing the codebook.
    original = load_workbook(template_path, read_only=True)
    original_rules = original['Rule Types']
    for row in original_rules.iter_rows(min_row=2):
        for cell in row:
            rules.cell(cell.row, cell.column).value = cell.value
    original.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
